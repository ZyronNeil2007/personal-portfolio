from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin   = Inches(1.2)
    section.right_margin  = Inches(1.2)

# ── Helpers ───────────────────────────────────────────────────────────────────
CYAN   = RGBColor(6, 182, 212)
WHITE  = RGBColor(255, 255, 255)
GRAY   = RGBColor(148, 163, 184)
LIME   = RGBColor(163, 230, 53)
DARK   = RGBColor(15, 23, 42)

def heading(text, level=1, color=None):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if color:
        for run in p.runs:
            run.font.color.rgb = RGBColor(*color)
    return p

def para(text, bold=False, italic=False, size=11, color=None, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = alignment
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p

def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        b = p.add_run(bold_prefix + ": ")
        b.bold = True
    p.add_run(text)
    return p

def divider():
    doc.add_paragraph()
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '06B6D4')
    pBdr.append(bottom)
    pPr.append(pBdr)

# ── Cover Page ────────────────────────────────────────────────────────────────
doc.add_paragraph()
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = t.add_run("Zyron Portfolio V2")
run.bold = True
run.font.size = Pt(34)
run.font.color.rgb = CYAN

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run2 = sub.add_run("Zyron Neil Bautista  |  CS Student & Creative Developer")
run2.font.size = Pt(13)
run2.font.color.rgb = GRAY

stack = doc.add_paragraph()
stack.alignment = WD_ALIGN_PARAGRAPH.CENTER
run3 = stack.add_run("React  |  Vite  |  TailwindCSS  |  GSAP  |  AOS  |  Matter.js")
run3.font.size = Pt(10)
run3.italic = True

datep = doc.add_paragraph()
datep.alignment = WD_ALIGN_PARAGRAPH.CENTER
datep.add_run(f"Generated: {datetime.date.today().strftime('%B %d, %Y')}").font.size = Pt(10)

doc.add_page_break()

# ── 1. HERO SECTION ───────────────────────────────────────────────────────────
heading("1. Hero Section (Home)", level=1)
para("Tagline / Eyebrow Badge:", bold=True, size=11)
para('  CS Student & Creative Developer  (with animated pulsing dot indicator)', size=11)

doc.add_paragraph()
para("Main Heading (3-line stagger animation via GSAP):", bold=True, size=11)
bullet("Line 1: \"Hi, I'm\"")
bullet("Line 2: \"Zyron Neil.\"  (bold white gradient)")
bullet("Line 3: \"I blend code with design.\"  (lime-yellow gradient)")

doc.add_paragraph()
para("Description:", bold=True, size=11)
para(
    '"I build functional, scalable web systems while making them visually engaging. '
    'Technology should not only work well — it should feel extraordinary."',
    italic=True, size=11
)

doc.add_paragraph()
para("Action Buttons:", bold=True, size=11)
bullet("View My Work  →  scrolls to #projects")
bullet("Let's Talk  →  scrolls to #contact")

doc.add_paragraph()
para("Stats Row:", bold=True, size=11)
for s in [("3+", "Projects Built"), ("2+", "Years Learning"), ("5+", "Tech Stacks")]:
    bullet(f"{s[1]}", bold_prefix=s[0])

doc.add_paragraph()
para("Right-side Portrait Card:", bold=True, size=11)
bullet("Image: neil_god.png  (full portrait, objectFit: cover)")
bullet("Floating badge (bottom-left): Stack — React · PHP · MySQL")
bullet("Floating badge (top-right): Focus — UI / UX Design")
bullet("Mouse-reveal mask effect powered by GSAP quick setters")

divider()

# ── 2. ABOUT SECTION ─────────────────────────────────────────────────────────
heading("2. About Me", level=1)

heading("Who I Am", level=2)
para(
    "A passionate tech enthusiast pursuing a degree in Computer Science. He enjoys turning "
    "complex ideas into real, functional systems — whether it's a web application, a database-driven "
    "dashboard, or a creative UI. Outside coding he enjoys design layouting, storytelling, and digital "
    "publications. His ultimate goal is to evolve as a versatile Software Engineer crafting premium "
    "interfaces that combine speed with visual elegance.",
    size=11
)

heading("What I Do", level=2)
for item in ["Develop web systems", "Create visual layouts", "Solve logic problems"]:
    bullet(item)

heading("Technical Arsenal", level=2)
categories = [
    ("Languages",           [("Java", "75%"), ("Python", "80%"), ("C++", "65%")]),
    ("Web Development",     [("HTML / CSS", "85%"), ("JavaScript", "75%"), ("Tailwind CSS", "90%")]),
    ("Databases & Tools",   [("MySQL / phpMyAdmin", "70%"), ("VS Code", "90%"), ("Git / GitHub", "80%")]),
    ("Design",              [("Layout Design", "85%"), ("Branding & Pubmats", "80%"), ("UI Design", "75%")]),
]
for cat_name, skills in categories:
    para(cat_name, bold=True, size=11)
    for skill, pct in skills:
        bullet(pct, bold_prefix=skill)

heading("Behind the Code (Fun Facts)", level=2)
for fact in [
    "Loves coffee & chill coding  ☕",
    "Plays Minecraft  🎮",
    "Enjoys movies, anime & BL  🎞️",
    "Detail-oriented & Creative  🎯",
]:
    bullet(fact)

divider()

# ── 3. JOURNEY SECTION ───────────────────────────────────────────────────────
heading("3. My Journey  (Timeline)", level=1)
para('"From Passion to Profession" — Every skill started as curiosity.', italic=True, size=11)

timeline = [
    (
        "2026 – Present  |  Dream in Progress",
        "Now a Computer Science student, Neil is pursuing Python & data mining, Java OOP, "
        "and full web development. Everything learned — from creativity to coding — is converging. "
        "Becoming a software engineer, web developer, and app developer is no longer just a dream.",
        ["Python", "Java OOP", "Web Dev", "Data Mining"],
    ),
    (
        "2025  |  Expanding Technical Skills",
        "Explored Computer Engineering concepts and worked with C++, deepening understanding of "
        "programming logic and systems. Challenged by Calculus 1 but built discipline and problem-solving skills.",
        ["C++", "Computer Engineering", "Problem Solving"],
    ),
    (
        "2024  |  Real-World Developer Era",
        "Went deeper into .NET Programming NC III — HTML, CSS, JavaScript, GitHub, and databases. "
        "Developed multiple web-based systems deployed at Barucboc National High School, "
        "gaining real-world experience with actual end users.",
        [".NET NC III", "JavaScript", "GitHub", "Databases"],
    ),
    (
        "2023  |  Peak Creativity & Programming Breakthrough",
        "Prime TikTok editing era with strong engagement. Simultaneously took Java Programming NC II "
        "and developed multiple systems in Senior High School, discovering a love for coding. "
        "Was introduced to HTML & CSS — instantly loved it for combining programming with design. "
        "Started self-studying and taking development more seriously.",
        ["Java NC II", "HTML & CSS", "TikTok Creator"],
    ),
    (
        "2022  |  Creative Growth",
        "Editing skills improved and developed abilities as a layout artist — posters and pubmats. "
        "Began to understand how design communicates ideas effectively.",
        ["Layout Artist", "Pubmats", "Visual Design"],
    ),
    (
        "2020–2021  |  Creative Exploration",
        "Started creating anime edits using CapCut and Alight Motion, exploring storytelling through "
        "visuals, music, and timing. This is where the creative foundation was built.",
        ["CapCut", "Alight Motion", "Storytelling"],
    ),
    (
        "2019  |  The Spark — Where It All Began",
        "Programming journey started with Minecraft. Running servers, configuring plugins, and customizing "
        "gameplay sparked curiosity about how code works behind the scenes. This was the beginning of everything.",
        ["Minecraft Servers", "Plugins", "Curiosity"],
    ),
]

for title_text, body, tags in timeline:
    heading(title_text, level=2)
    para(body, size=11)
    para("Tags: " + "  ·  ".join(tags), italic=True, size=10)

divider()

# ── 4. PROJECTS SECTION ──────────────────────────────────────────────────────
heading("4. Selected Work  (Projects)", level=1)
para("Filterable by: All · Web Dev · Layout & UI · Logic · Research · Business · Presentation", italic=True, size=10)
doc.add_paragraph()

projects = [
    (
        "Project 1  —  BNHS Online Quiz Website",
        "Web Dev · Logic",
        "A comprehensive online testing system featuring admin and student portals, complete with quiz management, "
        "interactive timers, results tracking, and an integrated feedback system.",
        "This system was developed to help Grade 11 students at Barucboc National High School access and answer "
        "review quizzes online. It features real-time scoring, database storage of student records, and interactive "
        "performance charts for teachers.",
        ["Web Dev", "phpMyAdmin", "Logic Design", "MySQL", "JavaScript"],
    ),
    (
        "Project 2  —  School Platform",
        "Web Dev · Layout & UI",
        "Designed and developed a fully functional school website tailored for seamless user experience, "
        "administrative resource sharing, and accessibility.",
        "A clean, modern platform designed for student informational needs. Includes glassmorphism cards, "
        "responsive menus, and light/dark theme adaptations to support accessible browsing.",
        ["HTML", "CSS", "JS", "UI/UX", "Mobile Responsive"],
    ),
    (
        "Project 3  —  Nephricarn Business",
        "Business · Layout & UI",
        "A creative food business concept offering premium burgers and healthy guava juice, tailored for "
        "student affordability, branding, and strategic market positioning.",
        "Developed during high school business simulation. Focuses on local ingredient sourcing, optimized "
        "pricing structures, and student-focused social media marketing layouts.",
        ["#Business", "#Entrepreneurship", "#FoodService", "#Marketing"],
    ),
    (
        "Project 4  —  BNHS Study: Online Quiz Research",
        "Research",
        'An academic research study titled "Effectiveness of Online Quiz Website to Grade 11 Students". '
        "Evaluated impacts on student motivation, accessibility, and grades.",
        "An empirical study that evaluated quantitative learning outcomes and engagement metrics of Grade 11 "
        "students before and after deploying the BNHS Online Quiz System. Showed a 22% improvement in retention scores.",
        ["#Research", "#Academic", "#DataAnalysis", "#Education"],
    ),
    (
        "Project 5  —  Layout & Design Portfolio",
        "Layout & UI",
        "A curated collection of visual publications, custom school posters, event pubmats, and vector logos "
        "focused on visual hierarchy, balance, and branding.",
        "Includes selected vector layouts, school publications, and award-winning posters designed for student "
        "organizations and graphic design contests.",
        ["#GraphicDesign", "#LayoutArtist", "#Branding", "#Creative"],
    ),
    (
        "Project 6  —  Research & Business Presentations",
        "Presentation",
        "Structured slides designed for professional academic defenses and investor business proposals, "
        "emphasizing typographic legibility and impactful slides.",
        "Highly visual presentation slides built using key slide design principles: high contrast, zero-clutter "
        "layouts, graphic metaphors, and structured information grouping.",
        ["#Presentation", "#VisualDesign", "#Communication", "#Academic"],
    ),
]

for p_title, p_cat, p_desc, p_details, p_tech in projects:
    heading(p_title, level=2)
    para("Category: " + p_cat, bold=True, size=10)
    para(p_desc, size=11)
    para("Deep Dive & Impact:", bold=True, size=10)
    para(p_details, italic=True, size=10)
    para("Technologies / Tags: " + "  ·  ".join(p_tech), size=10)

divider()

# ── 5. VISUAL DESIGN GALLERY ─────────────────────────────────────────────────
heading("5. Visual Design Gallery", level=1)
para("A masonry-layout gallery with hover sheen + clickable lightbox overlay:", size=11)
gallery = [
    ("image/quiz.png",        "BNHS Online Quiz Cover Mockup"),
    ("image/NOT_FINAL.jpg",   "Creative Visual Design Layout"),
    ("image/get_ready.jpg",   "Branding Poster Campaign"),
    ("image/me_booth.jpg",    "Zyron Neil — Portrait Event Capture"),
]
for src, alt in gallery:
    bullet(alt, bold_prefix=src)

divider()

# ── 6. CONTACT SECTION ───────────────────────────────────────────────────────
heading("6. Contact", level=1)
para(
    '"Let\'s build something great." — Always open to learning new things, collaborating on complex '
    'web systems, and building beautiful visual architectures.',
    italic=True, size=11
)
doc.add_paragraph()
for label, value in [
    ("Email",     "zyronneilbautista10@gmail.com"),
    ("Facebook",  "https://www.facebook.com/share/18ZFsaeo4S/"),
    ("Instagram", "https://www.instagram.com/zyronnei10/"),
    ("TikTok",    "https://www.tiktok.com/@zyron_neil"),
    ("GitHub",    "https://github.com/ZyronNeil2007"),
]:
    bullet(value, bold_prefix=label)

divider()

# ── 7. FOOTER ────────────────────────────────────────────────────────────────
heading("7. Footer", level=1)
para("© Designed & Developed by Zyron Neil Bautista", size=11)
para("Logo: image/logo.png", size=10)

divider()

# ── 8. TECH STACK NOTE ───────────────────────────────────────────────────────
heading("8. Technology Stack  (V2 React)", level=1)
tech_stack = [
    ("Framework",    "React 18 + Vite"),
    ("Styling",      "Tailwind CSS v3"),
    ("Animations",   "GSAP (GreenSock) — hero entrance, ambient glow pulse, mouse-mask reveal"),
    ("Scroll Anim",  "AOS (Animate On Scroll)"),
    ("Physics",      "Matter.js — Anti-Gravity Easter Egg (triple-click logo)"),
    ("Icons",        "Phosphor Icons (CDN)"),
    ("Language",     "JavaScript (JSX)"),
    ("Package Mgr",  "npm / node_modules"),
]
for label, value in tech_stack:
    bullet(value, bold_prefix=label)

# ── Save ──────────────────────────────────────────────────────────────────────
out = r"c:\Users\hp\Desktop\Web DEV Projects\zyron-portfolio-v2 react\Zyron_Portfolio_V2_Contents.docx"
doc.save(out)
print("Saved:", out)
